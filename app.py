import io
import json
import mimetypes
import os
import uuid
import smtplib
import ssl
import time
from urllib.error import URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from functools import lru_cache
from email.message import EmailMessage
from datetime import datetime
from flask import Flask, abort, render_template, request, redirect, url_for, session, flash, jsonify, Response, send_file
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from werkzeug.utils import secure_filename

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None


def load_environment_file(path=None):
    if path is None:
        path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.env')

    if not os.path.exists(path):
        return False

    with open(path, 'r', encoding='utf-8') as env_file:
        for raw_line in env_file:
            line = raw_line.strip()
            if not line or line.startswith('#') or '=' not in line:
                continue

            key, value = line.split('=', 1)
            key = key.strip().strip('"\'')
            value = value.strip().strip('"\'')
            os.environ[key] = value

    return True


if load_dotenv:
    load_dotenv()
else:
    load_environment_file()

app = Flask(__name__)
app.secret_key = os.getenv('FLASK_SECRET_KEY')
if not app.secret_key:
    raise RuntimeError('Set FLASK_SECRET_KEY in the ignored .env file before starting the app.')
UPLOAD_FOLDER = os.path.join(app.root_path, 'static', 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
DATA_FILE = os.getenv('PORTAL_STATE_FILE', os.path.join(app.root_path, 'database', 'portal_state.json'))

DEFAULT_USERS_DB = {}
DEFAULT_COMPLAINTS_DB = []
DEFAULT_EVIDENCE_FILES = []

DEFAULT_DEPARTMENTS = ["Academic Affairs", "IT Services", "Finance", "Student Housing"]
DEFAULT_COMPLAINT_CATEGORIES = ["Academic", "ICT", "Facilities / Housing", "Finance"]


def persist_state(users_data=None, complaints_data=None, departments_data=None, categories_data=None, evidence_data=None):
    if app.config.get('TESTING'):
        return

    os.makedirs(os.path.dirname(DATA_FILE), exist_ok=True)
    if users_data is None:
        users_data = globals().get('users_db', DEFAULT_USERS_DB)
    if complaints_data is None:
        complaints_data = globals().get('complaints_db', DEFAULT_COMPLAINTS_DB)
    if departments_data is None:
        departments_data = globals().get('departments', DEFAULT_DEPARTMENTS)
    if categories_data is None:
        categories_data = globals().get('complaint_categories', DEFAULT_COMPLAINT_CATEGORIES)
    if evidence_data is None:
        evidence_data = globals().get('evidence_files', DEFAULT_EVIDENCE_FILES)

    state = {
        'users_db': users_data,
        'complaints_db': complaints_data,
        'departments': departments_data,
        'complaint_categories': categories_data,
        'evidence_files': evidence_data,
    }
    with open(DATA_FILE, 'w', encoding='utf-8') as file:
        json.dump(state, file, indent=2)


def load_state():
    os.makedirs(os.path.dirname(DATA_FILE), exist_ok=True)
    if not os.path.exists(DATA_FILE):
        default_users = DEFAULT_USERS_DB.copy()
        default_complaints = DEFAULT_COMPLAINTS_DB.copy()
        default_departments = DEFAULT_DEPARTMENTS.copy()
        default_categories = DEFAULT_COMPLAINT_CATEGORIES.copy()
        default_evidence_files = DEFAULT_EVIDENCE_FILES.copy()
        persist_state(default_users, default_complaints, default_departments, default_categories, default_evidence_files)
        return default_users, default_complaints, default_departments, default_categories, default_evidence_files

    with open(DATA_FILE, 'r', encoding='utf-8') as file:
        try:
            state = json.load(file) or {}
        except json.JSONDecodeError:
            state = {}

    users = state.get('users_db') or DEFAULT_USERS_DB
    complaints = state.get('complaints_db') or DEFAULT_COMPLAINTS_DB
    departments_list = state.get('departments') or DEFAULT_DEPARTMENTS
    categories_list = state.get('complaint_categories') or DEFAULT_COMPLAINT_CATEGORIES
    evidence_files_list = state.get('evidence_files') or DEFAULT_EVIDENCE_FILES

    return users, complaints, departments_list, categories_list, evidence_files_list


def normalize_category(category):
    if not category:
        return 'Academic'

    cleaned = category.strip()
    if not cleaned:
        return 'Academic'

    lookup = {
        'it / network': 'ICT',
        'it network': 'ICT',
        'ict': 'ICT',
        'finance / nsfas': 'Finance',
        'finance': 'Finance',
        'nsfas': 'Finance',
        'nsfas / finance': 'Finance',
        'housing': 'Facilities / Housing',
        'facilities / housing': 'Facilities / Housing',
        'facilities housing': 'Facilities / Housing',
        'student housing': 'Facilities / Housing',
        'general': 'Academic',
    }

    return lookup.get(cleaned.lower(), cleaned)


def get_department_for_category(category):
    normalized = normalize_category(category)
    department_map = {
        'Academic': 'Academic Affairs',
        'ICT': 'IT Services',
        'Finance': 'Finance',
        'Facilities / Housing': 'Student Housing',
    }
    return department_map.get(normalized, 'Academic Affairs')


CHATBOT_MODEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ML_model', 'complaint_classifier.pkl')
CHATBOT_CATEGORY_MAP = {
    'Academics': 'Academic',
    'ICT Services': 'ICT',
    'Financial Aid': 'Finance',
    'Residences': 'Facilities / Housing',
}
CHATBOT_SAFETY_LABELS = {'Harassment & Protection', 'Campus Security'}
CHATBOT_EMERGENCY_CONTACT = 'Campus Security: 035 902 6000'
CHATBOT_CONFIDENCE_THRESHOLD = 0.4
CHATBOT_SAFETY_KEYWORDS = (
    'harassment', 'harassed', 'bullying', 'bullied', 'threatened', 'threatening',
    'stalking', 'stalked', 'assault', 'gender-based violence', 'gender based violence',
    'gbv', 'unsafe', 'campus security', 'stolen', 'theft', 'robbery',
)
CHATBOT_CATEGORY_KEYWORDS = {
    'ICT': ('wifi', 'wi-fi', 'internet', 'network', 'portal', 'login', 'log in', 'password', 'moodle', 'student email', 'computer lab', 'ict'),
    'Finance': ('allowance', 'nsfas', 'bursary', 'funding', 'financial aid', 'tuition', 'fees', 'fee', 'payment', 'refund', 'account balance', 'finance'),
    'Facilities / Housing': ('residence', 'hostel', 'housing', 'room', 'water', 'electricity', 'plumbing', 'maintenance', 'broken window', 'repair'),
    'Academic': ('exam', 'mark', 'marks', 'grade', 'results', 'transcript', 'lecturer', 'module', 'registration', 'timetable', 'remark', 'academic'),
}
CHATBOT_SYSTEM_PROMPT = (
    'You are the UNIZULU student grievance assistant. Respond warmly, briefly, and clearly in at most two sentences. '
    'Help students understand grievance submission and university support options. '
    'Do not claim to submit a grievance, access student records, or check its status; the portal handles those actions separately. '
    'Do not request passwords, student numbers, or unnecessary sensitive details. '
    'If a suggested grievance category is provided, explain it as a suggestion, not a final decision. '
    f'For immediate danger, tell the student to call {CHATBOT_EMERGENCY_CONTACT} or emergency services immediately.'
)


@lru_cache(maxsize=1)
def load_chatbot_model():
    import joblib

    return joblib.load(CHATBOT_MODEL_PATH)


def classify_chatbot_message(message):
    model = load_chatbot_model()
    probabilities = model.predict_proba([message])[0]
    labels = getattr(model, 'classes_', None)
    if labels is None:
        labels = model[-1].classes_

    best_index = max(range(len(probabilities)), key=probabilities.__getitem__)
    predicted_label = str(labels[best_index])
    confidence = float(probabilities[best_index])

    if predicted_label in CHATBOT_SAFETY_LABELS and confidence >= CHATBOT_CONFIDENCE_THRESHOLD:
        return {
            'category': None,
            'confidence': round(confidence, 4),
            'needs_safety_review': True,
        }

    category = CHATBOT_CATEGORY_MAP.get(predicted_label)
    if confidence < CHATBOT_CONFIDENCE_THRESHOLD or category is None:
        return {
            'category': None,
            'confidence': round(confidence, 4),
            'needs_safety_review': False,
        }

    return {
        'category': category,
        'confidence': round(confidence, 4),
        'needs_safety_review': False,
    }


def infer_chatbot_fallback(message):
    normalized = message.casefold()
    if any(keyword in normalized for keyword in CHATBOT_SAFETY_KEYWORDS):
        return {'category': None, 'confidence': None, 'needs_safety_review': True}

    matches = [
        category
        for category, keywords in CHATBOT_CATEGORY_KEYWORDS.items()
        if any(keyword in normalized for keyword in keywords)
    ]
    if len(matches) == 1:
        return {'category': matches[0], 'confidence': None, 'needs_safety_review': False}
    return None


def generate_chatbot_reply(messages, suggested_category=None):
    instructions = CHATBOT_SYSTEM_PROMPT
    if suggested_category:
        instructions += (
            f' The local classifier suggests {suggested_category}. Acknowledge the concern and explain '
            'that the portal will next ask for supporting evidence before submission.'
        )

    request_body = {
        'model': os.getenv('OLLAMA_MODEL', 'llama3.2:1b'),
        'messages': [{'role': 'system', 'content': instructions}, *messages],
        'stream': False,
        'keep_alive': '10m',
        'options': {'temperature': 0.2, 'num_ctx': 768, 'num_predict': 128},
    }
    base_url = os.getenv('OLLAMA_BASE_URL', 'http://127.0.0.1:11434').rstrip('/')
    ollama_request = Request(
        f'{base_url}/api/chat',
        data=json.dumps(request_body).encode('utf-8'),
        headers={'Content-Type': 'application/json'},
        method='POST',
    )
    try:
        with urlopen(ollama_request, timeout=120) as response:
            result = json.loads(response.read().decode('utf-8'))
    except URLError as error:
        raise RuntimeError('The local Ollama server is unavailable.') from error

    reply = result.get('message', {}).get('content', '').strip()
    if not reply:
        raise RuntimeError('The local model returned an empty response.')
    return reply


@lru_cache(maxsize=128)
def search_unizulu_information(query, cache_window):
    del cache_window
    params = urlencode({
        'search': query,
        'per_page': 3,
        '_fields': 'id,title,url,subtype',
    })
    search_request = Request(
        f'https://www.unizulu.ac.za/wp-json/wp/v2/search?{params}',
        headers={'Accept': 'application/json'},
    )
    try:
        with urlopen(search_request, timeout=5) as response:
            results = json.loads(response.read().decode('utf-8'))
    except (URLError, OSError, ValueError) as error:
        raise RuntimeError('Official UNIZULU search is temporarily unavailable.') from error

    if not isinstance(results, list):
        raise RuntimeError('Official UNIZULU search returned an unexpected response.')

    pages = []
    for result in results:
        if not isinstance(result, dict):
            continue
        url = result.get('url', '')
        if not isinstance(url, str) or not url.startswith('https://www.unizulu.ac.za/'):
            continue
        title = result.get('title')
        title = title.get('rendered') if isinstance(title, dict) else title
        if isinstance(title, str) and title.strip():
            pages.append({'title': title.strip(), 'url': url})

    return pages


@app.route('/chatbot/unizulu-search', methods=['POST'])
def chatbot_unizulu_search():
    if 'username' not in session or session.get('role') != 'student':
        return jsonify({'error': 'Please log in as a student to use the grievance assistant.'}), 401

    data = request.get_json(silent=True)
    query = data.get('query', '').strip() if isinstance(data, dict) and isinstance(data.get('query'), str) else ''
    if not 3 <= len(query) <= 200:
        return jsonify({'error': 'Search text must contain 3 to 200 characters.'}), 400

    try:
        pages = search_unizulu_information(query, int(time.time() // 300))
    except RuntimeError:
        app.logger.exception('Official UNIZULU search failed')
        return jsonify({'error': 'Official UNIZULU information is temporarily unavailable. Please try again shortly.'}), 503

    if not pages:
        return jsonify({
            'reply': 'I could not find a matching page on the official UNIZULU website. Try different keywords or contact the relevant university office.',
            'results': [],
        })

    links = '\n'.join(f"- {page['title']}: {page['url']}" for page in pages)
    return jsonify({
        'reply': f'Here are matching pages from the official UNIZULU website:\n{links}',
        'results': pages,
    })


def get_staff_department_filter():
    selected_department = session.get('department')
    if selected_department:
        return selected_department
    user = users_db.get(session.get('username'), {})
    return user.get('department', 'Academic Affairs')


def filter_complaints_for_staff():
    department = get_staff_department_filter()
    filtered = []

    for complaint in complaints_db:
        complaint_department = complaint.get('department') or get_department_for_category(complaint.get('category'))
        if complaint_department == department:
            filtered.append(complaint)

    return filtered


def get_database_connection():
    database_config = {
        'host': os.getenv('MYSQL_HOST'),
        'port': int(os.getenv('MYSQL_PORT', '3306')),
        'user': os.getenv('MYSQL_USER'),
        'password': os.getenv('MYSQL_PASSWORD', ''),
        'database': os.getenv('MYSQL_DATABASE'),
    }
    if not database_config['host'] or not database_config['user'] or not database_config['database']:
        raise RuntimeError('Set MYSQL_HOST, MYSQL_USER, and MYSQL_DATABASE in .env before submitting evidence.')

    try:
        import mysql.connector
    except ImportError as error:
        raise RuntimeError('Install the project requirements to enable MySQL evidence storage.') from error

    return mysql.connector.connect(**database_config)


def migrate_legacy_evidence():
    if not os.path.isdir(UPLOAD_FOLDER):
        return

    legacy_files = [
        filename for filename in os.listdir(UPLOAD_FOLDER)
        if os.path.isfile(os.path.join(UPLOAD_FOLDER, filename))
    ]
    if not legacy_files:
        return

    connection = get_database_connection()
    cursor = connection.cursor()
    migrated_files = []
    try:
        for stored_filename in legacy_files:
            file_path = os.path.join(UPLOAD_FOLDER, stored_filename)
            metadata = next((item for item in evidence_files if item.get('stored_filename') == stored_filename), None)
            with open(file_path, 'rb') as evidence_file:
                file_data = evidence_file.read()

            original_filename = (metadata or {}).get('original_filename')
            if not original_filename:
                original_filename = stored_filename.partition('_')[2] or stored_filename
            mime_type = (metadata or {}).get('mime_type') or mimetypes.guess_type(original_filename)[0] or 'application/octet-stream'
            complaint_id = (metadata or {}).get('complaint_id')
            cursor.execute(
                'INSERT INTO evidence_files '
                '(complaint_id, original_filename, stored_filename, stored_path, mime_type, size_bytes, file_data) '
                'VALUES (%s, %s, %s, %s, %s, %s, %s) '
                'ON DUPLICATE KEY UPDATE stored_path = VALUES(stored_path), mime_type = VALUES(mime_type), '
                'size_bytes = VALUES(size_bytes), file_data = VALUES(file_data)',
                (complaint_id, original_filename, stored_filename, f'database/{stored_filename}', mime_type, len(file_data), file_data),
            )
            migrated_files.append((file_path, metadata, stored_filename, len(file_data)))

        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()
        connection.close()

    for file_path, metadata, stored_filename, file_size in migrated_files:
        os.remove(file_path)
        if metadata:
            metadata['stored_path'] = f'database/{stored_filename}'
            complaint = next((item for item in complaints_db if item.get('id') == metadata.get('complaint_id')), None)
            if complaint:
                complaint['evidence_path'] = metadata['stored_path']
        else:
            evidence_files.append({
                'id': max((item.get('id', 0) for item in evidence_files), default=0) + 1,
                'complaint_id': None,
                'original_filename': stored_filename.partition('_')[2] or stored_filename,
                'stored_filename': stored_filename,
                'stored_path': f'database/{stored_filename}',
                'mime_type': mimetypes.guess_type(stored_filename)[0] or 'application/octet-stream',
                'size_bytes': file_size,
                'uploaded_at': datetime.now().isoformat(timespec='seconds'),
            })

    persist_state()


def save_uploaded_evidence(file_storage, complaint_id):
    if not file_storage or not file_storage.filename:
        return None

    original_filename = os.path.basename(file_storage.filename)
    safe_filename = secure_filename(original_filename) or 'evidence'
    unique_name = f"{uuid.uuid4().hex}_{safe_filename}"
    file_data = file_storage.read()
    metadata = {
        'id': max((item.get('id', 0) for item in evidence_files), default=0) + 1,
        'complaint_id': complaint_id,
        'original_filename': original_filename,
        'stored_filename': unique_name,
        'stored_path': f'database/{unique_name}',
        'mime_type': file_storage.mimetype or 'application/octet-stream',
        'size_bytes': len(file_data),
        'uploaded_at': datetime.now().isoformat(timespec='seconds'),
    }

    connection = get_database_connection()
    cursor = connection.cursor()
    try:
        cursor.execute(
            'INSERT INTO evidence_files '
            '(complaint_id, original_filename, stored_filename, stored_path, mime_type, size_bytes, file_data) '
            'VALUES (%s, %s, %s, %s, %s, %s, %s)',
            (complaint_id, original_filename, unique_name, metadata['stored_path'], metadata['mime_type'], len(file_data), file_data),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()
        connection.close()

    return metadata


def get_evidence_preview_url(evidence_path):
    if not evidence_path:
        return None
    return url_for('serve_evidence', stored_filename=os.path.basename(evidence_path))


app.jinja_env.globals['get_evidence_preview_url'] = get_evidence_preview_url


@app.route('/evidence/<path:stored_filename>')
def serve_evidence(stored_filename):
    if 'username' not in session:
        abort(401)

    try:
        migrate_legacy_evidence()
        connection = get_database_connection()
        cursor = connection.cursor()
        try:
            cursor.execute(
                'SELECT file_data, original_filename, mime_type, complaint_id '
                'FROM evidence_files WHERE stored_filename = %s',
                (os.path.basename(stored_filename),),
            )
            evidence = cursor.fetchone()
        finally:
            cursor.close()
            connection.close()
    except Exception:
        app.logger.exception('Unable to retrieve evidence from MySQL')
        abort(503)

    if not evidence or evidence[0] is None:
        abort(404)

    complaint = next((item for item in complaints_db if item.get('id') == evidence[3]), None)
    role = session.get('role')
    if role == 'student':
        allowed = complaint is not None and complaint.get('username') == session.get('username')
    elif role == 'staff':
        complaint_department = (complaint or {}).get('department') or get_department_for_category((complaint or {}).get('category'))
        allowed = complaint is not None and complaint_department == get_staff_department_filter()
    else:
        allowed = role == 'admin'

    if not allowed:
        abort(403)

    return send_file(
        io.BytesIO(evidence[0]),
        mimetype=evidence[2],
        download_name=evidence[1],
        as_attachment=False,
    )


def generate_reference_number():
    count = max((int(item.get('id', 0)) for item in complaints_db), default=0) + 1
    today = datetime.now().strftime('%Y%m%d')
    return f'GRV-{today}-{count:04d}'


def build_admin_analytics():
    category_counts = {}
    status_counts = {"Pending": 0, "In-Progress": 0, "Rejected": 0, "Resolved": 0, "Withdrawn": 0}

    for complaint in complaints_db:
        category = complaint.get('category', 'Academic')
        category_counts[category] = category_counts.get(category, 0) + 1

        status = complaint.get('status', 'Pending')
        if status == 'In Progress':
            status = 'In-Progress'
        if status in status_counts:
            status_counts[status] += 1
        else:
            status_counts[status] = status_counts.get(status, 0) + 1

    total_complaints = len(complaints_db) or 1
    category_percentages = {
        key: round((value / total_complaints) * 100, 1)
        for key, value in category_counts.items()
    }
    status_percentages = {
        key: round((value / total_complaints) * 100, 1)
        for key, value in status_counts.items()
    }

    return {
        'total_complaints': len(complaints_db),
        'status_counts': status_counts,
        'status_percentages': status_percentages,
        'category_counts': category_counts,
        'category_percentages': category_percentages,
        'latest_report': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    }

# In-memory portal state loaded from the configured data file
system_roles = ["student", "staff", "admin"]
users_db, complaints_db, departments, complaint_categories, evidence_files = load_state()
online_users = {}

EMAIL_HOST = os.getenv('EMAIL_HOST', 'smtp.gmail.com')
EMAIL_PORT = int(os.getenv('EMAIL_PORT', '587'))
EMAIL_USERNAME = (os.getenv('EMAIL_USERNAME') or '').strip() or None
EMAIL_PASSWORD = (os.getenv('EMAIL_PASSWORD') or '').strip() or None
EMAIL_FROM = os.getenv('EMAIL_FROM', 'no-reply@gmail.com')
EMAIL_USE_TLS = str(os.getenv('EMAIL_USE_TLS', 'true')).lower() == 'true'
EMAIL_USE_SSL = str(os.getenv('EMAIL_USE_SSL', 'false')).lower() == 'true'
EMAIL_TIMEOUT = float(os.getenv('EMAIL_TIMEOUT', '10'))
EMAIL_AUTH_METHOD = os.getenv(
    'EMAIL_AUTH_METHOD',
    'oauth2' if EMAIL_HOST.casefold() in {'smtp.outlook.com', 'smtp-mail.outlook.com'} else 'password',
).strip().casefold()
EMAIL_OAUTH_CLIENT_ID = (os.getenv('EMAIL_OAUTH_CLIENT_ID') or '').strip() or None
EMAIL_OAUTH_REFRESH_TOKEN = (os.getenv('EMAIL_OAUTH_REFRESH_TOKEN') or '').strip() or None
EMAIL_OAUTH_TENANT = (os.getenv('EMAIL_OAUTH_TENANT') or 'consumers').strip()


def get_outlook_access_token():
    token_url = f'https://login.microsoftonline.com/{EMAIL_OAUTH_TENANT}/oauth2/v2.0/token'
    token_request = Request(
        token_url,
        data=urlencode({
            'client_id': EMAIL_OAUTH_CLIENT_ID,
            'grant_type': 'refresh_token',
            'refresh_token': EMAIL_OAUTH_REFRESH_TOKEN,
            'scope': 'https://outlook.office.com/SMTP.Send offline_access',
        }).encode('ascii'),
        headers={'Content-Type': 'application/x-www-form-urlencoded'},
        method='POST',
    )
    try:
        with urlopen(token_request, timeout=EMAIL_TIMEOUT) as response:
            token_data = json.loads(response.read().decode('utf-8'))
    except (URLError, OSError, ValueError) as error:
        raise RuntimeError('Unable to obtain an Outlook access token.') from error

    access_token = token_data.get('access_token')
    if not isinstance(access_token, str) or not access_token:
        raise RuntimeError('Microsoft did not return an Outlook access token.')
    return access_token


def get_email_status():
    missing = []
    for key, value in {
        'EMAIL_HOST': EMAIL_HOST,
        'EMAIL_PORT': EMAIL_PORT,
        'EMAIL_USERNAME': EMAIL_USERNAME,
        'EMAIL_FROM': EMAIL_FROM,
    }.items():
        if key == 'EMAIL_PORT':
            continue
        if not value:
            missing.append(key)

    if EMAIL_AUTH_METHOD == 'password' and not EMAIL_PASSWORD:
        missing.append('EMAIL_PASSWORD')
    elif EMAIL_AUTH_METHOD == 'oauth2':
        if not EMAIL_OAUTH_CLIENT_ID:
            missing.append('EMAIL_OAUTH_CLIENT_ID')
        if not EMAIL_OAUTH_REFRESH_TOKEN:
            missing.append('EMAIL_OAUTH_REFRESH_TOKEN')
        if not EMAIL_OAUTH_TENANT:
            missing.append('EMAIL_OAUTH_TENANT')
    elif EMAIL_AUTH_METHOD != 'password':
        missing.append('EMAIL_AUTH_METHOD')

    if not 1 <= EMAIL_PORT <= 65535:
        missing.append('EMAIL_PORT')
    if EMAIL_TIMEOUT <= 0:
        missing.append('EMAIL_TIMEOUT')
    if EMAIL_USE_TLS and EMAIL_USE_SSL:
        missing.append('EMAIL_TLS_CONFIGURATION')

    return {
        'configured': not missing,
        'missing': missing,
        'host': EMAIL_HOST,
        'from_address': EMAIL_FROM,
        'use_tls': EMAIL_USE_TLS,
        'use_ssl': EMAIL_USE_SSL,
        'timeout': EMAIL_TIMEOUT,
        'auth_method': EMAIL_AUTH_METHOD,
    }


def get_user_email(username):
    user = users_db.get(username, {})
    return user.get('email') or None


def is_valid_email_address(address):
    if not isinstance(address, str) or len(address) > 254 or any(char.isspace() for char in address):
        return False

    if address.count('@') != 1:
        return False
    local_part, domain = address.rsplit('@', 1)
    if not local_part or len(local_part) > 64 or not domain or '.' not in domain:
        return False
    if local_part.startswith('.') or local_part.endswith('.') or '..' in local_part:
        return False

    labels = domain.split('.')
    return all(
        label
        and len(label) <= 63
        and label[0].isalnum()
        and label[-1].isalnum()
        and all(char.isalnum() or char == '-' for char in label)
        for label in labels
    )


def send_email_notification(subject, body, recipients):
    recipients = list(dict.fromkeys(email for email in recipients if email))
    if not recipients:
        app.logger.warning('Email notification skipped because it has no recipients.')
        return False
    status = get_email_status()
    if not status['configured']:
        app.logger.error('Email notification skipped; mail configuration is incomplete or invalid: %s', ', '.join(status['missing']))
        return False

    message = EmailMessage()
    message['Subject'] = subject
    message['From'] = EMAIL_FROM
    message['To'] = ', '.join(recipients)
    message.set_content(body)

    try:
        access_token = get_outlook_access_token() if EMAIL_AUTH_METHOD == 'oauth2' else None
        if EMAIL_USE_SSL:
            server_context = smtplib.SMTP_SSL(
                EMAIL_HOST,
                EMAIL_PORT,
                timeout=EMAIL_TIMEOUT,
                context=ssl.create_default_context(),
            )
        else:
            server_context = smtplib.SMTP(EMAIL_HOST, EMAIL_PORT, timeout=EMAIL_TIMEOUT)

        with server_context as server:
            if EMAIL_USE_TLS:
                server.starttls(context=ssl.create_default_context())
            if access_token:
                auth_string = f'user={EMAIL_USERNAME}\x01auth=Bearer {access_token}\x01\x01'
                server.auth('XOAUTH2', lambda _challenge: auth_string)
            else:
                server.login(EMAIL_USERNAME, EMAIL_PASSWORD)
            refused_recipients = server.send_message(message)
            if refused_recipients:
                app.logger.error(
                    'SMTP refused %d recipient(s); response codes: %s.',
                    len(refused_recipients),
                    ', '.join(sorted({str(result[0]) for result in refused_recipients.values()})),
                )
                return False
        return True
    except RuntimeError:
        app.logger.error('Outlook OAuth access-token acquisition failed; rerun the OAuth setup and verify Microsoft consent.')
        return False
    except smtplib.SMTPAuthenticationError as exc:
        if EMAIL_AUTH_METHOD == 'oauth2':
            app.logger.error(
                'Outlook OAuth authentication was rejected (code %s); rerun the OAuth setup and verify delegated SMTP.Send consent.',
                exc.smtp_code,
            )
        else:
            app.logger.error(
                'SMTP authentication was rejected (code %s); check EMAIL_USERNAME and EMAIL_PASSWORD.',
                exc.smtp_code,
            )
        return False
    except smtplib.SMTPRecipientsRefused as exc:
        response_codes = sorted({str(result[0]) for result in exc.recipients.values()})
        app.logger.error(
            'SMTP rejected all recipients (response codes: %s); verify recipient addresses and provider relay policy.',
            ', '.join(response_codes),
        )
        return False
    except smtplib.SMTPSenderRefused as exc:
        app.logger.error(
            'SMTP rejected the sender (code %s); verify EMAIL_FROM is permitted by the provider.',
            exc.smtp_code,
        )
        return False
    except (smtplib.SMTPException, OSError, ValueError) as exc:
        app.logger.error('Email delivery failed (%s).', type(exc).__name__)
        return False


def send_student_confirmation_email(complaint):
    student_email = get_user_email(complaint.get('username'))
    if not student_email:
        return False

    subject = f'Confirmation of complaint submission: {complaint.get("reference_number")}'
    body = (
        f"Hello {complaint.get('full_name')},\n\n"
        f"Your grievance has been submitted successfully.\n"
        f"Reference Number: {complaint.get('reference_number')}\n"
        f"Category: {complaint.get('category')}\n"
        f"Status: {complaint.get('status')}\n"
        f"Description: {complaint.get('description')}\n\n"
        f"Our staff will review your complaint and provide updates as needed."
    )
    return send_email_notification(subject, body, [student_email])


def send_complaint_notification(complaint):
    staff_emails = [
        get_user_email(username)
        for username, data in users_db.items()
        if data.get('role') in {'staff', 'admin'}
    ]

    staff_subject = f'New grievance submitted: {complaint.get("reference_number")}'
    staff_body = (
        f"Hello,\n\nA new grievance has been submitted.\n"
        f"Reference Number: {complaint.get('reference_number')}\n"
        f"Student: {complaint.get('full_name')}\n"
        f"Category: {complaint.get('category')}\n"
        f"Status: {complaint.get('status')}\n"
        f"Description: {complaint.get('description')}\n\n"
        f"Please log in to the grievance portal to review or update it."
    )

    staff_sent = send_email_notification(staff_subject, staff_body, [email for email in staff_emails if email])
    confirmation_sent = send_student_confirmation_email(complaint)
    if not staff_sent:
        app.logger.warning('New-grievance staff notification was not delivered.')
    return confirmation_sent


def send_status_update_notification(complaint, old_status, new_status):
    student_email = get_user_email(complaint.get('username'))
    if not student_email:
        return False

    subject = f'Grievance status update: {complaint.get("reference_number")}'
    body = (
        f"Hello {complaint.get('full_name')},\n\n"
        f"Your grievance status has been updated.\n"
        f"Reference Number: {complaint.get('reference_number')}\n"
        f"Previous Status: {old_status}\n"
        f"New Status: {new_status}\n"
        f"Category: {complaint.get('category')}\n\n"
        f"Please log in to the grievance portal to view the latest update."
    )
    return send_email_notification(subject, body, [student_email])


@app.before_request
def track_online_users():
    username = session.get('username')
    if not username:
        return

    user_data = users_db.get(username, {})
    online_users[username] = {
        'username': username,
        'full_name': session.get('full_name') or user_data.get('full_name', username),
        'role': session.get('role') or user_data.get('role', 'student'),
        'last_seen': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    }


def get_online_users():
    active_users = []
    for username, data in online_users.items():
        if not data:
            continue
        active_users.append({
            'username': data.get('username', username),
            'full_name': data.get('full_name') or users_db.get(username, {}).get('full_name', username),
            'role': data.get('role') or users_db.get(username, {}).get('role', 'student'),
            'last_seen': data.get('last_seen', '')
        })

    return sorted(active_users, key=lambda user: (user['role'], user['full_name']))


# Landing / Welcome Page
@app.route('/')
def index_page():
    return render_template('index.html')

@app.route('/home')
def home_page():
    return redirect(url_for('index_page'))


@app.route('/email_status')
def email_status():
    return jsonify(get_email_status())


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = (request.form.get('username') or '').strip()
        password = request.form.get('password') or ''
        role = request.form.get('role')
        department = (request.form.get('department') or '').strip()

        if not role:
            flash('Please select a role before logging in.', 'danger')
            return render_template('login.html', departments=departments)

        if role == 'student' and (not username.isdigit() or len(username) != 9):
            flash('Student number must be exactly 9 digits.', 'danger')
            return render_template('login.html', departments=departments)

        if role == 'staff' and not department:
            flash('Please select a department before logging in as staff.', 'danger')
            return render_template('login.html', departments=departments)

        user = users_db.get(username)
        if user and user['password'] == password and user['role'] == role:
            session['username'] = username
            session['full_name'] = user['full_name']
            session['role'] = user['role']
            if role == 'staff':
                session['department'] = department or user.get('department', 'Academic Affairs')
            else:
                session.pop('department', None)

            flash('Logged in successfully!', 'success')
            if role == 'staff':
                return redirect(url_for('staff_dashboard'))
            elif role == 'admin':
                return redirect(url_for('administrator'))
            else:
                return redirect(url_for('student_dashboard'))
        else:
            flash('Invalid username, password, or role selection.', 'danger')

    return render_template('login.html', departments=departments)

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = (request.form.get('username') or '').strip()
        full_name = (request.form.get('full_name') or '').strip()
        password = request.form.get('password') or ''
        role = request.form.get('role', 'student')
        email = (request.form.get('email') or '').strip()

        if not username or not full_name or not password:
            flash('Name, username, and password are required.', 'danger')
            return render_template('register.html')
        if not is_valid_email_address(email):
            flash('Please enter a valid email address for account updates.', 'danger')
            return render_template('register.html')
        if role == 'student' and (not username.isdigit() or len(username) != 9):
            flash('Student username must be exactly 9 digits.', 'danger')
            return render_template('register.html')

        if username in users_db:
            flash('Username already exists.', 'danger')
        else:
            users_db[username] = {
                'password': password,
                'full_name': full_name,
                'role': role,
                'email': email,
            }
            persist_state()
            flash('Registration successful! Please log in.', 'success')
            return redirect(url_for('login'))

    return render_template('register.html')

@app.route('/forgot_password')
@app.route('/forgot_password', methods=['POST'])
def forgot_password():
    if request.method == 'POST':
        username = (request.form.get('username') or '').strip()
        user = users_db.get(username)

        if (
            user
            and user.get('role') == 'student'
            and is_valid_email_address(user.get('email'))
        ):
            token_version = uuid.uuid4().hex
            user['password_reset_token_version'] = token_version
            persist_state()

            serializer = URLSafeTimedSerializer(app.secret_key, salt='password-reset')
            token = serializer.dumps({'username': username, 'version': token_version})
            reset_url = url_for('reset_password', token=token, _external=True)
            message = (
                'We received a request to reset your student portal password. '
                f'Use this link within one hour: {reset_url}\n\n'
                'If you did not request this reset, you can ignore this email.'
            )
            send_email_notification(
                'Student portal password reset',
                message,
                [user['email']],
            )

        flash(
            'If that student account is registered and has a valid email address on file, the system will attempt '
            'to send a reset link there. Check spam, and contact a system administrator if no email arrives.',
            'info',
        )
        return redirect(url_for('forgot_password'))

    return render_template('forgot_password.html')


@app.route('/reset_password/<token>', methods=['GET', 'POST'])
def reset_password(token):
    serializer = URLSafeTimedSerializer(app.secret_key, salt='password-reset')
    try:
        token_data = serializer.loads(token, max_age=3600)
    except SignatureExpired:
        flash('This reset link has expired. Please request another one.', 'danger')
        return redirect(url_for('forgot_password'))
    except BadSignature:
        flash('This reset link is invalid. Please request another one.', 'danger')
        return redirect(url_for('forgot_password'))

    username = token_data.get('username')
    user = users_db.get(username)
    if (
        not user
        or user.get('role') != 'student'
        or user.get('password_reset_token_version') != token_data.get('version')
    ):
        flash('This reset link is no longer valid. Please request another one.', 'danger')
        return redirect(url_for('forgot_password'))

    if request.method == 'POST':
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')
        if len(new_password) < 8:
            flash('Passwords must be at least 8 characters long.', 'danger')
        elif new_password != confirm_password:
            flash('The password confirmation does not match.', 'danger')
        else:
            user['password'] = new_password
            user['password_reset_token_version'] = uuid.uuid4().hex
            persist_state()
            flash('Your password has been reset. Please log in with your new password.', 'success')
            return redirect(url_for('login'))

    return render_template('reset_password.html')


@app.route('/student_dashboard')
def student_dashboard():
    if 'username' not in session or session.get('role') != 'student':
        flash('Please login as a student to access this page.', 'danger')
        return redirect(url_for('login'))

    student_complaints = [c for c in complaints_db if c['username'] == session['username']]
    latest_reference = next((c.get('reference_number') for c in reversed(student_complaints) if c.get('reference_number')), None)
    if latest_reference:
        session['last_reference_number'] = latest_reference

    return render_template('student_dashboard.html', complaints=student_complaints, latest_reference=latest_reference)


@app.route('/track_status')
def track_status():
    if 'username' not in session or session.get('role') != 'student':
        flash('Please login as a student to access this page.', 'danger')
        return redirect(url_for('login'))

    student_complaints = [c for c in complaints_db if c['username'] == session['username']]
    latest_reference = next((c.get('reference_number') for c in reversed(student_complaints) if c.get('reference_number')), None)
    if latest_reference:
        session['last_reference_number'] = latest_reference

    return render_template('track_status.html', complaints=student_complaints, latest_reference=latest_reference)


@app.route('/withdraw_grievance/<int:complaint_id>', methods=['POST'])
def withdraw_grievance(complaint_id):
    if 'username' not in session or session.get('role') != 'student':
        flash('Please login as a student to withdraw a grievance.', 'danger')
        return redirect(url_for('login'))

    complaint = next((c for c in complaints_db if c['id'] == complaint_id and c['username'] == session['username']), None)
    if complaint is None:
        flash('Grievance not found or it is not assigned to your account.', 'danger')
        return redirect(url_for('student_dashboard'))

    if complaint.get('status') in {'Resolved', 'Rejected', 'Withdrawn'}:
        flash('This grievance cannot be withdrawn because it is already resolved, rejected, or withdrawn.', 'danger')
        return redirect(url_for('student_dashboard'))

    complaint['status'] = 'Withdrawn'
    persist_state()
    flash(f'Grievance #{complaint_id} has been withdrawn successfully.', 'success')
    return redirect(url_for('student_dashboard'))


@app.route('/submit_grievance', methods=['POST'])
def submit_grievance():
    wants_json = request.accept_mimetypes.best == 'application/json'

    if 'username' not in session:
        if wants_json:
            return jsonify({'error': 'Please log in as a student to submit a grievance.'}), 401
        return redirect(url_for('login'))

    description = (request.form.get('description') or '').strip()
    category = normalize_category(request.form.get('category', 'Academic'))
    anonymous = request.form.get('anonymous') == 'on' or request.form.get('anonymous') == 'true'
    evidence_file = request.files.get('evidence')

    if not description:
        if wants_json:
            return jsonify({'error': 'Please provide a description for your grievance.'}), 400
        flash('Please provide a description for your grievance.', 'danger')
        return redirect(url_for('student_dashboard'))

    if not evidence_file or not evidence_file.filename:
        message = 'Evidence is required. Failure to submit supporting evidence may lead to rejection of your grievance.'
        if wants_json:
            return jsonify({'error': message}), 400
        flash(f'Grievance submission failed: {message}', 'danger')
        return redirect(url_for('student_dashboard'))

    new_id = max((int(item.get('id', 0)) for item in complaints_db), default=0) + 1
    try:
        migrate_legacy_evidence()
        evidence_metadata = save_uploaded_evidence(evidence_file, new_id)
    except Exception:
        app.logger.exception('Unable to store grievance evidence in MySQL')
        message = 'Evidence could not be saved. Please check that the MySQL database is configured and available, then try again.'
        if wants_json:
            return jsonify({'error': message}), 503
        flash(f'Grievance submission failed: {message}', 'danger')
        return redirect(url_for('student_dashboard'))

    reference_number = generate_reference_number()
    complaint = {
        "id": new_id,
        "full_name": 'Anonymous' if anonymous else session.get('full_name'),
        "username": session.get('username'),
        "description": description,
        "category": category,
        "department": get_department_for_category(category),
        "status": "Pending",
        "created_at": datetime.now().strftime('%Y-%m-%d'),
        "reference_number": reference_number,
        "is_anonymous": anonymous,
        "evidence_path": evidence_metadata['stored_path']
    }
    complaints_db.append(complaint)
    evidence_files.append(evidence_metadata)
    persist_state()

    email_sent = send_complaint_notification(complaint)

    session['last_reference_number'] = reference_number
    if wants_json:
        return jsonify({'reference_number': reference_number, 'email_sent': email_sent}), 201
    if not email_sent:
        flash('Your grievance was saved, but the confirmation email could not be sent. Contact support if you need help.', 'warning')
    flash(f'Grievance submitted successfully! Reference number: {reference_number}', 'success')
    return redirect(url_for('student_dashboard'))

def get_complaint_by_reference(reference_number):
    if 'username' not in session:
        return None

    return next((
        complaint for complaint in complaints_db
        if complaint.get('reference_number') == reference_number and complaint.get('username') == session.get('username')
    ), None)


@app.route('/chatbot/classify', methods=['POST'])
def chatbot_classify():
    if 'username' not in session or session.get('role') != 'student':
        return jsonify({'error': 'Please log in as a student to use the grievance assistant.'}), 401

    data = request.get_json(silent=True)
    if not isinstance(data, dict) or not isinstance(data.get('message'), str):
        return jsonify({'error': 'A text message is required.'}), 400

    message = data['message'].strip()
    if len(message) < 8:
        return jsonify({'error': 'Please describe the issue in a little more detail.'}), 400
    if len(message) > 2000:
        return jsonify({'error': 'Please keep the issue description under 2000 characters.'}), 400

    try:
        result = classify_chatbot_message(message)
    except (ImportError, OSError, ValueError) as error:
        app.logger.error('Local grievance classifier is unavailable: %s', error)
        return jsonify({'error': 'The local classifier is unavailable. Please use the grievance form.'}), 503

    return jsonify(result)


@app.route('/chatbot/chat', methods=['POST'])
def chatbot_chat():
    if 'username' not in session or session.get('role') != 'student':
        return jsonify({'error': 'Please log in as a student to use the grievance assistant.'}), 401

    data = request.get_json(silent=True)
    messages = data.get('messages') if isinstance(data, dict) else None
    if not isinstance(messages, list) or not messages or len(messages) > 12:
        return jsonify({'error': 'Send up to 12 recent chat messages.'}), 400

    clean_messages = []
    total_characters = 0
    for item in messages:
        if not isinstance(item, dict) or item.get('role') not in {'user', 'assistant'}:
            return jsonify({'error': 'Chat messages must have a user or assistant role.'}), 400
        content = item.get('content')
        if not isinstance(content, str) or not content.strip() or len(content) > 2000:
            return jsonify({'error': 'Each chat message must contain 1 to 2000 characters.'}), 400
        total_characters += len(content)
        clean_messages.append({'role': item['role'], 'content': content.strip()})

    if total_characters > 8000 or clean_messages[-1]['role'] != 'user':
        return jsonify({'error': 'Send a user message with no more than 8000 total characters.'}), 400

    latest_message = clean_messages[-1]['content']
    try:
        classification = classify_chatbot_message(latest_message)
        if not classification['category'] and not classification['needs_safety_review']:
            classification = infer_chatbot_fallback(latest_message) or classification
        if classification['needs_safety_review']:
            return jsonify({
                'reply': (
                    'This may involve harassment or campus safety. I cannot notify responders automatically. '
                    f'If anyone is in immediate danger, call {CHATBOT_EMERGENCY_CONTACT} or emergency services immediately.'
                ),
                **classification,
            })

        if classification['category']:
            reply = (
                f"This may fit the {classification['category']} category. "
                'I can help you submit it through the guided grievance flow.'
            )
        else:
            reply = generate_chatbot_reply(clean_messages)
    except Exception:
        app.logger.exception('Generative grievance assistant is unavailable')
        return jsonify({'error': 'The local generative assistant is unavailable. Start Ollama and ensure the configured model is installed.'}), 503

    return jsonify({'reply': reply, **classification})


@app.route('/complaint_status', methods=['POST'])
def complaint_status_post():
    if 'username' not in session:
        return jsonify({'error': 'Please log in to check complaint status.'}), 401

    data = request.get_json(silent=True) or {}
    reference_number = (data.get('reference_number') or request.form.get('reference_number') or '').strip()
    if not reference_number:
        return jsonify({'error': 'A reference number is required.'}), 400

    complaint = get_complaint_by_reference(reference_number)
    if complaint is None:
        return jsonify({'error': 'No complaint found for this reference number.'}), 404

    return jsonify({
        'reference_number': complaint.get('reference_number'),
        'category': complaint.get('category'),
        'status': complaint.get('status'),
        'description': complaint.get('description'),
        'created_at': complaint.get('created_at')
    })


@app.route('/complaint_status/<reference_number>')
def complaint_status(reference_number):
    if 'username' not in session:
        return jsonify({'error': 'Please log in to check complaint status.'}), 401

    complaint = get_complaint_by_reference(reference_number)
    if complaint is None:
        return jsonify({'error': 'No complaint found for this reference number.'}), 404

    return jsonify({
        'reference_number': complaint.get('reference_number'),
        'category': complaint.get('category'),
        'status': complaint.get('status'),
        'description': complaint.get('description'),
        'created_at': complaint.get('created_at')
    })


@app.route('/staff_dashboard')
def staff_dashboard():
    if 'username' not in session or session.get('role') != 'staff':
        flash('Access restricted to staff only.', 'danger')
        return redirect(url_for('login'))

    selected_department = session.get('department') or users_db.get(session.get('username'), {}).get('department', 'Academic Affairs')
    if not selected_department:
        flash('Please select a department before viewing grievances.', 'danger')
        return redirect(url_for('login'))

    staff_complaints = filter_complaints_for_staff()
    return render_template('staff_dashboard.html', complaints=staff_complaints, selected_department=selected_department, departments=departments)

@app.route('/update_status/<int:complaint_id>', methods=['POST'])
def update_status(complaint_id):
    if 'username' not in session or session.get('role') != 'staff':
        return redirect(url_for('login'))

    new_status = request.form.get('status')
    notification_sent = None
    for c in complaints_db:
        if c['id'] == complaint_id:
            previous_status = c.get('status')
            c['status'] = new_status
            persist_state()
            if previous_status != new_status:
                notification_sent = send_status_update_notification(c, previous_status, new_status)
            break

    flash(f'Grievance #{complaint_id} status updated to {new_status}.', 'success')
    if notification_sent is False:
        flash('The status was saved, but the student email could not be sent. Check the mail server configuration.', 'danger')
    return redirect(url_for('staff_dashboard'))

# Administrator dashboard route
@app.route('/administrator')
def administrator():
    if 'username' not in session or session.get('role') != 'admin':
        flash('Access restricted to administrators.', 'danger')
        return redirect(url_for('login'))

    users_list = [
        {
            "username": username,
            "full_name": data["full_name"],
            "role": data.get("role", "student"),
            "department": data.get("department", "General"),
            "email": data.get("email", ""),
        }
        for username, data in users_db.items()
    ]
    analytics = build_admin_analytics()
    return render_template(
        'Administrator.html',
        users_list=users_list,
        roles=system_roles,
        departments=departments,
        complaint_categories=complaint_categories,
        analytics=analytics,
        online_users_list=get_online_users()
    )


@app.route('/administrator/add_user', methods=['POST'])
def add_user():
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    username = request.form.get('username', '').strip()
    full_name = request.form.get('full_name', '').strip()
    email = (request.form.get('email') or '').strip()
    role = request.form.get('role', 'student')
    department = request.form.get('department', 'Academic Affairs')

    if not username or not full_name:
        flash('Username and full name are required.', 'danger')
    elif not is_valid_email_address(email):
        flash('Please enter a valid email address for the new user.', 'danger')
    elif username in users_db:
        flash(f'User {username} already exists.', 'danger')
    else:
        users_db[username] = {
            'password': request.form.get('password', '123'),
            'full_name': full_name,
            'role': role,
            'department': department,
            'email': email,
        }
        persist_state()
        flash(f'User {full_name} added successfully.', 'success')

    return redirect(url_for('administrator'))


@app.route('/administrator/update_user_role/<username>', methods=['POST'])
def update_user_role(username):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    if username in users_db:
        email = (request.form.get('email') or '').strip()
        if not is_valid_email_address(email):
            flash('Please enter a valid email address for this user.', 'danger')
            return redirect(url_for('administrator'))
        users_db[username]['email'] = email
        users_db[username]['role'] = request.form.get('role', users_db[username].get('role', 'student'))
        users_db[username]['department'] = request.form.get('department', users_db[username].get('department', 'Academic Affairs'))
        persist_state()
        flash(f'User {username} updated successfully.', 'success')

    return redirect(url_for('administrator'))


@app.route('/administrator/reset_password/<username>', methods=['POST'])
def reset_user_password(username):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    user = users_db.get(username)
    new_password = request.form.get('new_password', '')
    confirm_password = request.form.get('confirm_password', '')

    if user is None:
        flash(f'User {username} was not found.', 'danger')
    elif len(new_password) < 8:
        flash('Passwords must be at least 8 characters long.', 'danger')
    elif new_password != confirm_password:
        flash('The password confirmation does not match.', 'danger')
    else:
        user['password'] = new_password
        user['password_reset_token_version'] = uuid.uuid4().hex
        persist_state()
        flash(f'Password reset for {username}. Share the new password with the user directly.', 'success')

    return redirect(url_for('administrator'))


@app.route('/administrator/remove_user/<username>', methods=['POST'])
def remove_user(username):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    user = users_db.pop(username, None)
    if user:
        persist_state()
        flash(f'User {username} removed successfully.', 'success')
    else:
        flash(f'User {username} was not found.', 'danger')

    return redirect(url_for('administrator'))


@app.route('/administrator/add_department', methods=['POST'])
def add_department():
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    department_name = request.form.get('department', '').strip()
    if department_name and department_name not in departments:
        departments.append(department_name)
        persist_state()
        flash(f'Department {department_name} added successfully.', 'success')
    else:
        flash('Department name is invalid or already exists.', 'danger')

    return redirect(url_for('administrator'))


@app.route('/administrator/add_category', methods=['POST'])
def add_category():
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    category_name = request.form.get('category', '').strip()
    if category_name and category_name not in complaint_categories:
        complaint_categories.append(category_name)
        persist_state()
        flash(f'Complaint category {category_name} added successfully.', 'success')
    else:
        flash('Category name is invalid or already exists.', 'danger')

    return redirect(url_for('administrator'))


@app.route('/administrator/remove_department/<path:department>', methods=['POST'])
def remove_department(department):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    normalized = department.strip()
    if normalized in departments:
        departments.remove(normalized)
        persist_state()
        flash(f'Department {normalized} removed successfully.', 'success')
    else:
        flash(f'Department {normalized} was not found.', 'danger')

    return redirect(url_for('administrator'))


@app.route('/administrator/remove_category/<path:category>', methods=['POST'])
def remove_category(category):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    normalized = category.strip()
    if normalized in complaint_categories:
        complaint_categories.remove(normalized)
        persist_state()
        flash(f'Category {normalized} removed successfully.', 'success')
    else:
        flash(f'Category {normalized} was not found.', 'danger')

    return redirect(url_for('administrator'))


@app.route('/administrator/generate_report')
def generate_report():
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    analytics = build_admin_analytics()
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = 'Summary'
    summary_sheet['A1'] = 'Grievance Report'
    summary_sheet['A1'].font = Font(size=16, bold=True, color='FFFFFF')
    summary_sheet['A1'].fill = PatternFill('solid', fgColor='003366')
    summary_sheet.merge_cells('A1:C1')
    summary_sheet['A1'].alignment = Alignment(horizontal='center')

    summary_sheet['A3'] = 'Generated At'
    summary_sheet['A3'].font = Font(bold=True)
    summary_sheet['B3'] = analytics['latest_report']
    summary_sheet['A5'] = 'Total Complaints'
    summary_sheet['A5'].font = Font(bold=True)
    summary_sheet['B5'] = analytics['total_complaints']
    summary_sheet['B5'].font = Font(bold=True)

    header_fill = PatternFill('solid', fgColor='D9EAF7')
    summary_sheet['A7'] = 'Status'
    summary_sheet['A7'].font = Font(bold=True)
    summary_sheet['A7'].fill = header_fill
    summary_sheet['B7'] = 'Count'
    summary_sheet['B7'].font = Font(bold=True)
    summary_sheet['B7'].fill = header_fill
    summary_sheet['C7'] = 'Percentage'
    summary_sheet['C7'].font = Font(bold=True)
    summary_sheet['C7'].fill = header_fill
    status_row = 8
    for status in ['Pending', 'In-Progress', 'Rejected', 'Resolved', 'Withdrawn']:
        summary_sheet.cell(row=status_row, column=1, value=status)
        summary_sheet.cell(row=status_row, column=2, value=analytics['status_counts'].get(status, 0))
        summary_sheet.cell(row=status_row, column=3, value=f"{analytics['status_percentages'].get(status, 0)}%")
        status_row += 1

    light_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for cell in summary_sheet['A7:C' + str(status_row - 1)]:
        for item in cell:
            item.border = light_border

    status_chart = BarChart()
    status_chart.title = 'Status Breakdown'
    status_chart.y_axis.title = 'Count'
    status_chart.x_axis.title = 'Status'
    data = Reference(summary_sheet, min_col=2, min_row=7, max_row=11, max_col=2)
    categories = Reference(summary_sheet, min_col=1, min_row=8, max_row=11)
    status_chart.add_data(data, titles_from_data=False)
    status_chart.set_categories(categories)
    status_chart.height = 7
    status_chart.width = 13
    summary_sheet.add_chart(status_chart, 'E7')

    category_sheet = workbook.create_sheet('Category Breakdown')
    category_sheet['A1'] = 'Category'
    category_sheet['B1'] = 'Count'
    category_sheet['C1'] = 'Percentage'
    category_sheet['A1'].font = Font(bold=True)
    category_sheet['B1'].font = Font(bold=True)
    category_sheet['C1'].font = Font(bold=True)
    category_sheet['A1'].fill = header_fill
    category_sheet['B1'].fill = header_fill
    category_sheet['C1'].fill = header_fill
    category_row = 2
    for category, count in analytics['category_counts'].items():
        category_sheet.cell(row=category_row, column=1, value=category)
        category_sheet.cell(row=category_row, column=2, value=count)
        category_sheet.cell(row=category_row, column=3, value=f"{analytics['category_percentages'].get(category, 0)}%")
        category_row += 1

    for cell in category_sheet['A1:C' + str(category_row - 1)]:
        for item in cell:
            item.border = light_border

    pie_chart = PieChart()
    pie_chart.title = 'Complaint Categories'
    pie_data = Reference(category_sheet, min_col=2, min_row=1, max_row=category_row - 1, max_col=2)
    pie_categories = Reference(category_sheet, min_col=1, min_row=2, max_row=category_row - 1, max_col=1)
    pie_chart.add_data(pie_data, titles_from_data=False)
    pie_chart.set_categories(pie_categories)
    pie_chart.height = 7
    pie_chart.width = 12
    category_sheet.add_chart(pie_chart, 'E4')

    detail_sheet = workbook.create_sheet('Detailed Complaints')
    detail_headers = ['Reference Number', 'Student Username', 'Student Name', 'Category', 'Department', 'Status', 'Created At', 'Anonymous']
    for column_index, header in enumerate(detail_headers, start=1):
        cell = detail_sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9EAF7')
        cell.alignment = Alignment(horizontal='center')

    for row_index, complaint in enumerate(complaints_db, start=2):
        detail_sheet.cell(row=row_index, column=1, value=complaint.get('reference_number', ''))
        detail_sheet.cell(row=row_index, column=2, value=complaint.get('username', ''))
        detail_sheet.cell(row=row_index, column=3, value=complaint.get('full_name', ''))
        detail_sheet.cell(row=row_index, column=4, value=complaint.get('category', ''))
        detail_sheet.cell(row=row_index, column=5, value=complaint.get('department', get_department_for_category(complaint.get('category'))))
        detail_sheet.cell(row=row_index, column=6, value=complaint.get('status', 'Pending'))
        detail_sheet.cell(row=row_index, column=7, value=complaint.get('created_at', ''))
        detail_sheet.cell(row=row_index, column=8, value='Yes' if complaint.get('is_anonymous') else 'No')

    for row in detail_sheet.iter_rows(min_row=1, max_row=detail_sheet.max_row, min_col=1, max_col=8):
        for cell in row:
            cell.border = light_border

    for sheet in [summary_sheet, category_sheet, detail_sheet]:
        for column_cells in sheet.columns:
            visible_cells = [cell for cell in column_cells if hasattr(cell, 'column_letter')]
            if not visible_cells:
                continue
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in visible_cells)
            sheet.column_dimensions[visible_cells[0].column_letter].width = min(max_length + 2, 26)

    summary_sheet.freeze_panes = 'A8'
    category_sheet.freeze_panes = 'A2'
    detail_sheet.freeze_panes = 'A2'

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = Response(output.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response.headers['Content-Disposition'] = 'attachment; filename=grievance_report.xlsx'
    return response


@app.route('/logout')
def logout():
    username = session.get('username')
    if username:
        online_users.pop(username, None)
    session.clear()
    flash('Logged out successfully.', 'info')
    return redirect(url_for('login'))

if __name__ == '__main__':
    app.run(debug=True)